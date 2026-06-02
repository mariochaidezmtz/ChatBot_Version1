"""
Módulo para cargar y procesar documentos múltiples
Objetivo: Extraer texto de múltiples documentos y dividirlos en chunks manejables
Soporta: PDF, TXT, MD, DOCX
"""

from langchain_text_splitters import RecursiveCharacterTextSplitter
from typing import List
from langchain_core.documents import Document
from pathlib import Path
import os


class DocumentProcessor:
    def __init__(self, chunk_size: int = 1000, chunk_overlap: int = 200):
        """
        Inicializa el procesador de documentos
        
        Args:
            chunk_size: Tamaño de cada chunk en caracteres
            chunk_overlap: Superposición entre chunks (para continuidad)
        """
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n\n", "\n", ".", " ", ""]
        )
    
    def load_file(self, file_path: str) -> str:
        """
        Carga un archivo y extrae el texto según su extensión.
        Soporta: .pdf, .txt, .md, .docx
        """
        path = Path(file_path)
        ext = path.suffix.lower()

        try:
            if ext == ".pdf":
                from pypdf import PdfReader
                reader = PdfReader(file_path)
                text = ""
                print(f"Leyendo PDF: {file_path} ({len(reader.pages)} páginas)")
                for page in reader.pages:
                    text += page.extract_text() + "\n"

            elif ext in (".txt", ".md"):
                with open(file_path, "r", encoding="utf-8") as f:
                    text = f.read()
                print(f"Leyendo texto: {file_path}")

            elif ext == ".docx":
                import docx
                doc = docx.Document(file_path)
                text = "\n".join([p.text for p in doc.paragraphs])
                print(f"Leyendo DOCX: {file_path}")

            else:
                raise ValueError(f"Formato no soportado: {ext}")

            print(f"[OK] {path.name} cargado ({len(text)} caracteres)")
            return text

        except Exception as e:
            print(f"[ERROR] Error al cargar {file_path}: {e}")
            raise
        
    def split_into_chunks(self, text: str, source: str = "manual") -> List[Document]:
        """
        Divide el texto en chunks más pequeños
        
        Args:
            text: Texto a dividir
            source: Fuente del documento (para metadata)
            
        Returns:
            Lista de Document objects con metadata
        """
        chunks = self.splitter.split_text(text)
        
        documents = [
            Document(
                page_content=chunk,
                metadata={
                    "source": source,
                    "chunk_id": i,
                    "length": len(chunk)
                }
            )
            for i, chunk in enumerate(chunks)
        ]
        
        print(f"[INFO] Documento dividido en {len(documents)} chunks")
        print(f"   - Tamaño promedio: {sum(len(d.page_content) for d in documents) // len(documents)} caracteres")
        
        return documents


def scan_documents_directory(directory: str = "./documents", extensions: List[str] = None) -> List[str]:
    """
    Escanea dinámicamente una carpeta buscando archivos de documentos.
    
    Args:
        directory: Ruta de la carpeta a escanear
        extensions: Lista de extensiones a buscar (default: ['.pdf', '.txt', '.md', '.docx'])
    
    Returns:
        Lista de rutas de archivos encontrados
    """
    if extensions is None:
        extensions = ['.pdf', '.txt', '.md', '.docx']
    
    directory_path = Path(directory)
    
    if not directory_path.exists():
        print(f"[WARNING] Directorio no encontrado: {directory}")
        return []
    
    found_files = []
    
    for ext in extensions:
        # Buscar archivos con la extensión especificada
        for file_path in directory_path.glob(f"*{ext}"):
            found_files.append(str(file_path))
    
    # También buscar en subdirectorios recursivamente
    for ext in extensions:
        for file_path in directory_path.rglob(f"*{ext}"):
            if str(file_path) not in found_files:
                found_files.append(str(file_path))
    
    print(f"[INFO] Archivos encontrados en {directory}: {len(found_files)}")
    for file_path in found_files:
        print(f"   - {Path(file_path).name}")
    
    return found_files


def prepare_documents(paths=None, directory: str = None) -> List[Document]:
    """
    Carga y procesa uno o varios documentos.
    
    Args:
        paths: str (un solo archivo), list (varios archivos), o None para usar escaneo de directorio
        directory: Ruta del directorio a escanear si paths es None
    
    Returns:
        Lista de documentos procesados y divididos
    """
    processor = DocumentProcessor(chunk_size=1000, chunk_overlap=200)
    all_documents = []
    
    # Si no se proporcionan paths, escanear directorio
    if paths is None:
        if directory is None:
            directory = "./documents"
        paths = scan_documents_directory(directory)
        
        if not paths:
            print("[WARNING] No se encontraron documentos para procesar")
            return []
    
    # Convertir a lista si es un solo path
    if isinstance(paths, str):
        paths = [paths]

    for file_path in paths:
        try:
            text = processor.load_file(file_path)
            source_name = Path(file_path).name
            docs = processor.split_into_chunks(text, source=source_name)
            all_documents.extend(docs)
        except Exception as e:
            print(f"[ERROR] Error procesando {file_path}: {e}")
            continue

    print(f"[INFO] Total de chunks generados: {len(all_documents)}")
    return all_documents